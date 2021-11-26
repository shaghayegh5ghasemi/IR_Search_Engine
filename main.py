import openpyxl
from pathlib import Path
from hazm import *
import json
from datetime import datetime
import sys


class Document():
    def __init__(self, id, content, title):
        self.id = id
        self.content = content
        self.title = title


# this function collects the documents info from the excel file
def getDocuments():
    raw_doc_path = Path('IR1_7k_news.xlsx')
    raw_doc = openpyxl.load_workbook(raw_doc_path)
    doc_sheet = raw_doc.active
    content = doc_sheet['A']
    title = doc_sheet['C']
    documents = []  # extract contents and their title from .xlsx file
    for i in range(1, len(content)):
        documents.append(Document(i, content[i].value, title[i].value))
    return documents


# this function prepossesses the document: tokenize, normalize, stemming, delete stop words
def preprocess(raw_document):
    normalizer = Normalizer()
    stemmer = Stemmer()
    lemmatizer = Lemmatizer()
    stopwordslist = set(stopwords_list())
    stopwordslist.update(['.', ':', '?', '!', '/', '//', '*', '[', ']', '{', '}', ';', '\'', '\"', '(', ')', ''])
    tokens = []
    words = []
    words_temp = word_tokenize(normalizer.normalize(raw_document.content))
    for word in words_temp:
        if word not in stopwordslist:
            temp = lemmatizer.lemmatize(word)
            temp = stemmer.stem(temp)
            words.append(temp)
    for word in set(words):
        # my tokens consists of: term, positions in the document
        positions = [i + 1 for i, x in enumerate(words) if x == word]
        tokens.append([word, positions, len(positions)])
    return tokens

def positionalIndex(pos_idx, id, tokens):
    # t[0] = word, t[1] = list of position, t[2] = number of repeatitions
    for t in tokens:
        if t[0] not in pos_idx.keys():
            pos_idx[t[0]] = [t[2], [(id, t[1], t[2])]]
        else:
            pos_idx[t[0]][0] += t[2]
            pos_idx[t[0]][1].append((id, t[1], t[2]))


# check if the index is constructed before or not
def isConstructed():
    path_to_file = 'Positional_Index.json'
    path = Path(path_to_file)
    if path.is_file():
        return True
    return False

def retriev_doc(word, pos_idx):
    docIDs = []
    if word in pos_idx.keys():
        postings_list = pos_idx[word][1]
        for i in range(len(postings_list)):
            docIDs.append(postings_list[i][0])
    return docIDs, postings_list

def respond():
    pass

if __name__ == '__main__':
    documentCollection = getDocuments()
    if isConstructed():
        # load positional index
        with open('Positional_Index.json') as json_file:
            pos_index = json.load(json_file)
    else:
        pos_idx = {}  # our positional index
        for document in documentCollection:
            tokens = preprocess(document)
            positionalIndex(pos_idx, document.id, tokens)
        with open('Positional_Index.json', 'w') as convert_file:
            convert_file.write(json.dumps(pos_idx))  # save the index to a file

    normalizer = Normalizer()
    stemmer = Stemmer()
    lemmatizer = Lemmatizer()
    stopwordslist = set(stopwords_list())
    stopwordslist.update(['.', ':', '?', '!', '/', '//', '*', '[', ']', '{', '}', ';', '\'', '\"', '(', ')', ''])
    while (True):
        query = input("Query/Enter F to finish: ")
        start_time = datetime.now()
        if(query == 'f' or query == 'F'):
            sys.exit()
        query_words_temp = word_tokenize(normalizer.normalize(query))
        query_words = []
        for word in query_words_temp:
            if word not in stopwordslist:
                temp = lemmatizer.lemmatize(word)
                temp = stemmer.stem(temp)
                query_words.append(temp)
        if (len(query_words) == 1):
            docIDs, _ = retriev_doc(query_words[0], pos_index)
            print("")
            print("Documents which contain the word in ther query: ")
            print(docIDs)
            print("")
            print("Title of one related news: ")
            print(documentCollection[docIDs[0] - 1].title)
            print("")
        if (len(query_words) > 1):
            postings_lists = []
            for word in query_words:
                _ , posting = retriev_doc(word, pos_index)
                postings_lists.append(posting)
            # inja bayad eshterak begiram
        end_time = datetime.now()
        print('Duration: {}'.format(end_time - start_time))






